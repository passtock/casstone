from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / '03_연구계획_및_정리노트/02_연구_및_실험계획서'
SOURCE = FOLDER / '2026.09.08연구계획서.docx'
OUTPUT = FOLDER / '2026.09.08연구계획서_모델프롬프트피처추가.docx'
doc = Document(SOURCE)
anchor = next(p for p in doc.paragraphs if p.text.startswith('7. 연구 범위'))
anchor.paragraph_format.page_break_before = True

def font(run, size=10, bold=False, color='222222'):
    run.font.name = '맑은 고딕'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')

def place(element):
    anchor._p.addprevious(element)

def para(text, bold=False, size=10, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.widow_control = True
    font(p.add_run(text), size, bold)
    place(p._p)
    return p

def heading(text):
    p = para(text, True, 12, 9)
    p.style = doc.styles['Heading 2']
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    return p

def sub(text):
    p = para(text, True, 10, 5)
    p.paragraph_format.keep_with_next = True
    return p

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8EEF5' if ri == 0 else ('F7F9FB' if ri % 2 == 0 else 'FFFFFF'))
            pr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), '65')
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.03
                for r in p.runs:
                    font(r, 9, ri == 0)
    place(t._tbl)
    para('', after=2)
    return t

def code(text):
    p = para(text, size=8.8, after=7)
    p.paragraph_format.line_spacing = 1.0
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'F2F4F7')
    p._p.get_or_add_pPr().append(shade)
    return p

# Keep the original study intact; clarify the acquisition window required by the added features.
for t in doc.tables:
    for row in t.rows:
        if row.cells[0].text == '6. 종료':
            row.cells[1].text = '접촉 해제 여부를 관찰한다. 해제 지시 후 최소 2초까지 녹화를 유지한 뒤 종료한다.'
            row.cells[2].text = '고정 관찰창 동안 홈 복귀를 지시하지 않음. 이후 복귀는 채점에서 제외. 시간상한은 개발 단계에서 고정.'
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs: font(r, 9)
for p in doc.paragraphs:
    if p.text.startswith('출력: gradable 여부'):
        p.text = '출력: gradable 여부, grasp shape, hold, active release, frame ID 및 feature key에 연결한 근거, 제한점. 구체 프롬프트와 출력 규칙은 6.11–6.12에 제시한다.'
        for r in p.runs: font(r, 10.5)
    if p.text.startswith('작성일:'):
        p.add_run('  |  모델·입력·프롬프트 보완: 2026년 9월 9일')
        for r in p.runs: font(r, 8.5)

heading('6.8 사용할 모델과 비교 범위')
para('본 연구는 GPT-5.2를 주 평가 모델로 사전 지정하고, Gemini 2.5 Pro로 동일 입력 비교의 재현성을 확인한다. 모델 간 성능 순위를 찾는 것이 아니라, 각 모델 안에서 영상·운동학·건강인 참조를 추가했을 때 전문가 판단과의 일치도가 달라지는지를 검증한다. 아래 구성은 연구 실행 계획이며 본 과제에서의 성능 우위가 확인된 결과는 아니다.')
table(['구분', '모델과 식별자', '역할과 선택 근거'], [
    ['주 VLM', 'GPT-5.2\ngpt-5.2-2025-12-11', 'A1–A4와 주 가설 검정에 사용. 이미지 입력, 구조화 출력, 고정 스냅샷 지원을 선정 근거로 한다.[6, 7]'],
    ['보조 VLM', 'Gemini 2.5 Pro\ngemini-2.5-pro', '같은 환자·같은 입력 조건으로 보조 재현 분석. stable ID를 불변 스냅샷으로 간주하지 않고 실제 응답 버전·호출일을 기록한다.[8]'],
    ['수치 기준모델 A0', 'L2 규제 순서형 로지스틱 회귀', '0 < 1 < 2를 예측. A3 수치 블록을 동일하게 펼쳐 사용하며 환자 단위 LOSO로 학습·검증한다.'],
    ['선택적 로컬 후보', 'Qwen/Qwen2.5-VL-7B-Instruct', '외부 전송 제한이나 로컬 재현이 필요한 경우의 별도 탐색 후보. GPU 메모리·14장 입력·처리시간 확인 전 필수 비교에 포함하지 않는다.[9]']
], [2.5, 5.7, 9.3])
sub('동일 조건을 유지하는 실행 설정')
para('주모델은 reasoning.effort=medium을 초기 고정안으로 사용하고, 시행마다 새로운 독립 API 요청을 보낸다. 도구 호출, 웹검색, 이전 시행의 대화와 결과는 제공하지 않는다. temperature·seed는 해당 모델 및 추론 설정에서 지원이 확인된 경우에만 명시하며, 지원하지 않는 인자를 임의로 넣지 않는다. 보조 모델의 추론 설정은 별도로 기록하되 그 모델의 A1–A4 안에서는 동일하게 유지한다.')
para('입력 영상은 원본 RGB에서 추출한 14장의 개별 이미지로 제공한다. GPT-5.2에 원본 AVI를 직접 입력하지 않는다. 이미지 순서·해상도·ROI·시각 표시는 모든 영상 조건에서 동일하게 고정한다. 모델 교체 또는 API 지원 종료가 발생하면 시험 전에 계획을 갱신하며, 서로 다른 모델 버전의 응답을 한 조건의 결과로 합치지 않는다.')
para('동일 입력은 독립적으로 3회 실행한다. 첫 번째 계획된 실행을 주 분석으로 사용하고 나머지는 반복성 분석에만 사용한다. 보조 모델의 결과가 좋아도 이를 근거로 주모델을 사후 변경하지 않는다. 동일 환자에서 수행하는 공급자 간 재현 비교는 별도 기관·별도 환자의 외부검증과 구분한다.')

heading('6.9 VLM에 제공할 운동학 피처')
para('주 거리 신호는 a(t) = 1000 × ||Pthumb_RS(t) − Pindex_RS(t)||로 계산하며 단위는 mm이다. 좌표는 D455의 정렬 depth로 복원한 카메라 좌표를 사용한다. MediaPipe world 거리와 혼합하거나, RS 결측을 MP 값으로 자동 대체하지 않는다. MP 거리는 센서·모델 간 비교를 위한 별도 기록으로 보존한다.')
table(['피처 키', '정의와 단위', '해석 및 결측 처리'], [
    ['mga_mm', '도달 시작부터 최초 가시적 접촉 전까지 a(t)의 최대값, mm', '접촉 또는 도달 시작을 판독할 수 없거나 peak 주변 품질이 부족하면 null.'],
    ['tmga_ratio', '(t_MGA − t_reach) / (t_contact − t_reach), 0–1', '분모가 0 이하이거나 사건이 결측이면 null. 여러 동일 최대값은 가장 이른 시각.'],
    ['reach_duration_s', 't_contact − t_reach, s', 'TMGA의 구간 길이를 함께 제공. 빠르거나 느리다는 이유만으로 성공 판정하지 않음.'],
    ['preparation_time_s', '준비 버튼 시각 − 시작 지시 시각, s', '준비 버튼은 참가자의 준비 신호이지 목표 파지 성공 라벨이 아님.'],
    ['hold_aperture_sd_mm', '준비 신호 후 지정 2초 유지구간의 aperture 표준편차, mm', '유효 표본 2개 이상 및 품질 기준 충족 시 계산. 낮은 변동이 유지 성공을 뜻하지 않음.'],
    ['release_aperture_change_mm', '해제 후 1.8–2.0초 중앙값 − 해제 전 0.2초 중앙값, mm', '두 시간창이 모두 유효해야 계산. 양수는 거리 증가이며 능동 해제의 직접 증거가 아님.'],
    ['release_aperture_mm', '해제 지시 후 0, 0.5, 1.0, 1.5, 2.0초의 aperture 5점, mm', '각 점과 유효 마스크를 함께 제공. 허용 시간창 밖의 표본을 끌어오거나 외삽하지 않음.']
], [4.5, 6.4, 6.6])
para('해제 궤적은 각 목표 시각의 ±0.05초 안에 있는 유효값의 중앙값을 사용하는 초기안이다. 해제 전 기준창은 [−0.2, 0)초로 정의한다. 창 폭과 관측 가능성은 개발 건강인 4명에서 확인한 뒤 고정한다. 접촉 해제가 빨라도 해제 지시 후 2초까지 기록하고 그동안 홈 복귀를 지시하지 않아, 빠른 시행만 조기에 결측되는 편향을 방지한다.')
para('t_reach와 t_contact는 도달 시작·최초 가시적 접촉의 사건 시각이다. 최종 전문가 등급과 분리한 사건 주석 절차에서 생성하고, 판독 출처와 불확실성을 보존한다. 목표 파지 완성 시각을 최초 접촉 시각으로 대체하지 않는다. 검증되지 않은 관절각, TAM/TAROM, SPARC, 굴곡·신전 속도는 주 입력에서 제외한다.')

heading('6.10 품질 정보와 건강인 참조 입력')
table(['입력 블록', '제공할 값', '사용 원칙'], [
    ['measurement_quality', '구간별 valid_ratio(0–1), longest_gap_s, interpolated_ratio(0–1), peak_valid, contact_valid', '결측을 수행 실패로 읽지 않도록 품질을 피처와 함께 제공.'],
    ['missingness', '피처별 valid 및 reason\n예: depth_hole, depth_edge, event_unreadable, insufficient_samples', '값은 null, 이유는 별도 문자열. 실제 0과 구분하며 평균값으로 모델 입력을 채우지 않음.'],
    ['healthy_reference', '과제·손 조건, 실제 n, 평균, SD, 중앙값, IQR', '반복시행을 먼저 참가자 중앙값으로 요약. 가능한 최대 n=20이나 결측이 있으면 실제 n 기록.'],
    ['reference_z', '안정적인 피처의 (x − 평균) / SD', 'SD=0 또는 참조가 불안정하면 null. 정상/비정상 또는 성공/실패 임계값으로 쓰지 않음.'],
    ['공통 과제 정보', '과제 종류, 물체 치수, 지정 유지시간, 프레임 순서·시각', 'A1–A4에 동일하게 제공. 환자 여부·중증도·FMA·전문가 판독은 제외.']
], [3.4, 6.4, 7.7])
sub('개발 단계에서 고정할 품질 기준')
para('초기 후보는 구간별 valid_ratio ≥ 0.80, longest_gap ≤ 0.25초로 두고, MGA·접촉 전후 ±0.10초에서는 유효한 실제 관측을 요구한다. 0.10초 이하의 짧은 내부 공백만 보간 후보로 허용하되 peak·접촉 주변은 보간으로 유효성을 만들어내지 않는다. 기준 미달은 해당 피처의 null 처리로 연결한다. 이 수치는 검증된 임상 기준이 아니며 개발 자료에서 확정하고 시험 환자 결과로 조정하지 않는다.')
sub('건강인 참조와 A0의 피처 일치')
para('건강인 참조는 과제·우세손 조건을 포함해 사전에 정한 규칙으로 매칭한다. 마비측을 무조건 비우세손 참조에 대응시키지 않는다. 손 길이 정규화를 추가하려면 실측 손 길이와 같은 정의의 참조분포를 함께 확보하고 입력 명세를 갱신한다. 결측이 많거나 분산이 불안정한 참조 피처는 제외하고 n과 제외 사유를 보존한다.')
para('A0에는 A3의 scalar 피처, 해제 궤적 5점, 유효 마스크, 품질·참조 수치와 과제 one-hot을 고정 순서로 입력한다. frame ID와 시행 ID는 식별용이며 예측 피처에서 제외한다. 원시 수치와 Z-score의 중복·공선성을 고려해 규제를 적용하고, 피처를 결과에 따라 선별하지 않는다. 결측 대체·표준화·규제강도 선택이 필요하면 모두 LOSO의 training 환자 내부에서만 수행한다.')

heading('6.11 공통 프롬프트와 조건별 입력')
para('프롬프트는 한국어 zero-shot 지침을 기본으로 한다. 별도의 예시 정답을 포함하지 않아 개발 건강인에서 거의 2점인 예시가 응답을 편향시키는 위험을 줄인다. 과제별 형태 기준 문구는 전문가 평가표와 동일하게 정의하되, 환자 시험 전에 고정한 기준 설명만 제공한다.')
sub('시스템 지침 예시')
para('당신은 손 파지 연구의 수행 판독자이다. 제공된 자료만 사용하여 목표 파지 형태, 지정 2초 유지, 손가락 개방에 의한 능동적 접촉 해제를 각각 평가하라. 이것은 연구용 수행 판독이며 표준 FMA 채점이나 진단이 아니다. 준비 버튼과 종료 버튼은 성공의 증거가 아니다. aperture 증가만으로 접촉 해제 또는 능동성을 단정하지 말라. 건강인 참조는 수치 해석의 맥락이며 성공·실패 임계값이 아니다. 제공되지 않은 영상·수치와 프레임 사이 동작을 만들어내지 말라. 근거가 부족하면 unreadable을 사용하라. 각 판단에는 제공된 frame_id 또는 feature_key와 짧은 관찰 근거를 연결하라. 최종 0/1/2 등급을 직접 생성하지 말고 지정된 JSON만 반환하라.')
sub('사용자 메시지 템플릿')
code('trial_id: {익명 시행 ID}\n'
     'task: {cylinder | sphere | lateral_pinch}\n'
     'object_dimensions_mm: {과제 물체의 고정 치수}\n'
     'rubric: {목표 형태·2초 유지·능동 접촉 해제의 공통 기준}\n'
     'frames: {F01–F14 이미지, 순서, 상대 시각}\n'
     'kinematics: {피처 값, 단위, valid, 결측 이유, measurement_quality}\n'
     'healthy_reference: {과제·손 조건, n, 분포, 허용된 reference_z}\n'
     '요청: 제공된 블록만 근거로 각 항목을 판독하고 지정 JSON을 반환하라.')
table(['조건', '실제 제공 블록', '제외 블록'], [
    ['A1', '공통 과제 정보 + 원본 RGB 14장', 'kinematics, healthy_reference'],
    ['A2', 'A1 + kinematics 및 품질·결측 정보', 'healthy_reference'],
    ['A3', 'A2 + healthy_reference 및 reference_z', '없음'],
    ['A4', '공통 과제 정보 + A3의 수치·품질·참조', 'frames 전체와 이미지 첨부']
], [1.8, 9.5, 6.2])
para('조건 이름 A1–A4와 “더 많은 정보이므로 더 정확하다” 등의 유도 문구는 모델에 보내지 않는다. 제외 블록은 0으로 채우지 않고 요청에서 제거한다. A1–A3의 이미지는 수치나 골격이 덧그려지지 않은 원본 RGB로 통일한다. 저장된 MediaPipe 영상은 점검용이며 A1 입력으로 쓰지 않는다.')

heading('6.12 JSON 출력과 코드 기반 등급 결정')
para('세부 상태와 근거를 구조화해 반환받고, 최종 등급은 외부 코드가 결정한다. 아래는 입력 형식을 설명하기 위한 출력 예시이며 실제 환자 관측 결과가 아니다. 예시의 unreadable은 판독 실패를 허용하는 출력 형식을 보여준다.')
example = {
    'schema_version': '1.0', 'gradable': False,
    'grasp_shape': {'status': 'unreadable', 'frame_ids': [], 'feature_keys': [], 'evidence': '목표 형태를 판독할 근거가 부족함'},
    'hold': {'status': 'unreadable', 'frame_ids': [], 'feature_keys': [], 'evidence': '지정 유지 여부를 판독할 근거가 부족함'},
    'active_release': {'status': 'unreadable', 'frame_ids': [], 'feature_keys': [], 'evidence': '능동적 접촉 해제를 확인할 근거가 부족함'},
    'limitations': ['제공된 정보만으로 수행 등급을 결정할 수 없음']
}
code(json.dumps(example, ensure_ascii=False, indent=2))
para('실제 JSON Schema는 additionalProperties=false 및 필수 키를 사용한다. status는 grasp_shape의 경우 complete/incomplete/unreadable, hold·active_release의 경우 여기에 not_applicable을 추가한 enum으로 제한한다. 존재하지 않는 frame_id·feature_key, 조건상 제공하지 않은 근거, 모순된 상태 조합은 검증 실패로 기록한다. gradable은 아래 등급 결정 가능 여부와 일치하는지 후처리에서 확인한다.', size=9.3)
table(['출력 상태와 운영 기록', '코드의 최종 등급'], [
    ['통증 중단 등 사전 정의된 기준 NA, 또는 판단에 필요한 정보 부족', 'NA'],
    ['grasp_shape가 incomplete로 명확히 판독됨', '0; 이후 항목은 not_applicable 가능'],
    ['grasp_shape complete + hold 또는 active_release 중 하나 이상 incomplete', '1'],
    ['grasp_shape·hold·active_release가 모두 complete', '2']
], [12.6, 4.9])

heading('6.13 실행 기록과 결측을 고려한 비교')
sub('프롬프트와 입력을 고정하는 순서')
para('개발 건강인 4명에서 촬영 구도·14장 샘플링·2초 해제 관찰창·피처 품질 기준·프롬프트·출력 검증을 점검한다. 이 자료는 최종 환자 성능에 포함하지 않는다. 첫 시험 환자 전에 모델 ID, 이미지 크기와 ROI, 피처 명세, 참조분포 버전, 프롬프트와 JSON Schema를 고정하고 변경 이력을 남긴다. 현재 손 추적·저장 코드의 존재를 이 VLM 평가 파이프라인의 구현·검증 완료로 간주하지 않는다.')
table(['보존 항목', '기록 내용'], [
    ['입력 식별', '익명 participant_id/trial_id, 조건, 원본 frame ID·시각, 이미지/피처/참조 파일의 SHA256'],
    ['모델 실행', '모델 ID·응답 버전, 호출일, 추론 설정, 출력 토큰 한도, API 요청 ID, 지연시간과 사용량'],
    ['응답 및 후처리', '응답 원문, JSON 검증 결과, 세부 상태, 코드 산출 등급, 반복 번호 1–3'],
    ['오류', 'API 전송 실패, JSON/schema 실패, 근거 참조 실패, 판독 불가를 별도 분류']
], [3.2, 14.3])
para('전송 실패는 동일 요청으로 최대 2회까지 재시도하는 초기안으로 정하고 모든 시도를 남긴다. 형식이 잘못된 모델 응답은 유리한 결과가 나올 때까지 재질문하거나 사람이 수정하지 않는다. 독립 반복 3회의 첫 응답이 실패한 경우에도 2·3회차의 좋은 응답으로 주 결과를 바꾸지 않는다. 최대 225 trial × 4조건 × 3회 = 모델당 2,700회이며, 보조 VLM까지 모두 수행하면 최대 5,400회이다. 이는 예정 호출 수로 실제 실행 결과가 아니다.')
sub('NA와 응답 가능률을 분리한 분석')
para('전문가 기준 NA는 판독 가능한 참조 등급이 없는 시행이다. 모델 NA, API 실패 및 JSON 실패와 구분한다. 각 조건에서 기준 판독 가능 시행 중 숫자 등급을 낸 비율을 coverage로 보고한다. A4는 영상이 없기 때문에 파지 형태·능동 해제에 대한 판독 불가가 증가할 수 있으며, 이 자체도 영상 제거 효과로 보고한다.')
para('A3–A1의 MAE는 두 조건 모두 숫자 예측을 낸 공통 시행에서 짝지어 계산하고 coverage를 함께 제시한다. 기준 판독 가능 시행 전체에서 모델 NA·실패에 최대 절대오차 2를 부여한 민감도 분석도 사전 지정한다. 쉬운 시행만 남긴 낮은 MAE를 전체 성능 향상으로 해석하지 않는다. 환자의 세 과제 중 공통 유효 시행이 없는 과제가 있으면 주 환자 평균을 임의의 두 과제 평균으로 대체하지 않고, 제외·coverage와 민감도 분석에 별도로 보고한다.')
para('건강인 참조에는 환자 라벨이 들어가지 않는다. A0의 training/test 분리는 참가자 단위로 유지하고 같은 환자의 다른 시행·반대 손이 training에 섞이지 않게 한다. A0는 지도학습 기준모델이고 VLM은 환자 라벨로 학습하지 않은 조건이라는 학습 정보 차이를 결과 해석에 명시한다.')

# Add verified technical sources before the final freeze checklist.
references = [
    '[6] OpenAI. GPT-5.2 model documentation. 이미지 입력, 구조화 출력 및 스냅샷 식별자. https://developers.openai.com/api/docs/models/gpt-5.2 (확인 2026-09-09).',
    '[7] OpenAI. Structured model outputs. JSON Schema 기반 출력 제약. https://developers.openai.com/api/docs/guides/structured-outputs (확인 2026-09-09).',
    '[8] Google. Gemini 2.5 Pro model documentation. 이미지 입력, 구조화 출력 및 모델 코드. https://ai.google.dev/gemini-api/docs/models/gemini-2.5-pro (확인 2026-09-09).',
    '[9] Qwen. Qwen2.5-VL-7B-Instruct official model card. https://huggingface.co/Qwen/Qwen2.5-VL-7B-Instruct (확인 2026-09-09).'
]
ref_anchor = next(p for p in doc.paragraphs if p.text.startswith('University of Gothenburg'))
previous = ref_anchor._p
for text in references:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs: font(r, 9)
    previous.addnext(p._p)
    previous = p._p

# Mark the title semantically without changing the retained two-line subject.
doc.paragraphs[0].style = doc.styles['Title']
for r in doc.paragraphs[0].runs: font(r, 13, True, '000000')
for p in doc.paragraphs[:3]:
    for border in list(p._p.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
for border in list(doc.styles['Title'].element.xpath('./w:pPr/w:pBdr')):
    border.getparent().remove(border)
# Keep the retained interpretation table from leaving a single orphan row.
compact = False
for p in doc.paragraphs:
    if p.text.startswith('6.1 '): compact = True
    if p.text.startswith('6.8 '): compact = False
    if compact:
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.03
for t in doc.tables:
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn('w:cantSplit')) is None:
            trpr.append(OxmlElement('w:cantSplit'))
    if len(t.rows) > 1:
        header = t.rows[0]._tr.get_or_add_trPr()
        if header.find(qn('w:tblHeader')) is None:
            header.append(OxmlElement('w:tblHeader'))
    if t.cell(0, 0).text == '결과':
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs: font(r, 8.5, ri == 0, 'FFFFFF' if ri == 0 else '222222')
doc.core_properties.modified = __import__('datetime').datetime(2026, 9, 9, 0, 0, 0)
doc.save(OUTPUT)
print(str(OUTPUT))
