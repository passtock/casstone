# -*- coding: utf-8 -*-
"""0928 발표덱 -> 수정본 패치 스크립트 (재현용).

원본을 그대로 복사한 뒤 '덱 쪽에서 고칠 부분'만 최소 수정한다.
근거: outputs/03-검증/deck-0928-vs-plan-v6-alignment-audit.md
정본: outputs/research-plan-v6.md (v6.3)

수정 원칙:
- 레이아웃을 깨지 않도록 '문단 수를 늘리지 않는다'. 가시 텍스트는 기존 줄을 대체만 한다.
- 내용 추가는 전부 발표자 노트(단일 문단·단일 런)에만 넣는다.
- 원본 파일은 건드리지 않는다.
"""
import shutil
from pathlib import Path
from pptx import Presentation

SRC = Path("capstone/손과제_VLM_연구계획_0928.pptx")
DST = Path("capstone/손과제_VLM_연구계획_0928_수정본.pptx")

shutil.copyfile(SRC, DST)
prs = Presentation(str(DST))

TAG = " [수정본 2026-09-28 ↔ v6.3 대조]"


def notes_run(slide):
    """노트의 유일한 런을 반환(없으면 예외)."""
    paras = slide.notes_slide.notes_text_frame.paragraphs
    assert len(paras) == 1 and len(paras[0].runs) == 1, "노트 구조가 예상과 다름"
    return paras[0].runs[0]


def edit_notes(slide, old, new):
    r = notes_run(slide)
    assert old in r.text, f"치환 대상 없음: {old[:40]!r}"
    r.text = r.text.replace(old, new, 1)


def append_notes(slide, extra):
    r = notes_run(slide)
    r.text = r.text.rstrip() + " " + extra


# ── S4: Li 2026 인용 정정 (provenance) ────────────────────────────────
s4 = prs.slides[3]
s4.shapes[7].text_frame.paragraphs[0].runs[0].text = "FMA 부분집단 28명, 영상 899개"

edit_notes(
    s4,
    "Li 2026: https://journals.plos.org/digitalhealth/article?id=10.1371/journal.pdig.0001506 "
    "FMA Methods, Table 2, Fig 6. 4 healthy+24 stroke, 899 videos.",
    "Li 2026 (출판본: PLOS Digit Health 5(7):e0001506, PMID 42406872, "
    "https://doi.org/10.1371/journal.pdig.0001506 — 비장애 20 + 뇌졸중 51; "
    "프리프린트 arXiv:2511.17727 = 비장애 29 + 51로 cohort가 다름). "
    "FMA 실험 부분집단 4 healthy+24 stroke = 28명, 899 videos.",
)
append_notes(s4, TAG + " Li 2026/2022, Tang 2025 식별자·cohort를 1차 출처로 재확인함.")

# ── S6: D-15 실행 격차 명시 ───────────────────────────────────────────
s6 = prs.slides[5]
append_notes(
    s6,
    "실행 격차(D-15): 현재 파이프라인 make_conditions.py는 아직 K1/K2 대신 "
    "MGA_mm_3D_cal·PV_mm_s·SPARC·TAM_total_deg를, 과제도 free/cylinder를 사용한다. "
    "본평가 수집 전 정렬 필요(research-plan-v6.md §7.1)." + TAG,
)

# ── S9: 탐색적 조건 명시(주 비교 아님) ───────────────────────────────
s9 = prs.slides[8]
append_notes(
    s9,
    "탐색적(보정 없음, CI·효과크기만): A3−A4(영상의 가치), A0/A0-time, "
    "오류주입 bias_1u/3u·burst_3u 기전(H5). 이들은 PR1~PR3 주 비교가 아니다. "
    "오류 주입은 A2·A3에만 적용한다(research-plan-v6.md §6.1)." + TAG,
)

# ── S11: 1급 기술 결과(보류율·Pareto) 보강 — 정본 §0·§1.3·§9 ─────────
s11 = prs.slides[10]
append_notes(
    s11,
    "1급 기술 결과(검정 아님, 반드시 보고): 보류율(A3) + (보류율, MAE) Pareto 곡선"
    "(Q1 임계값 sweep, 컴퓨트 의존) + 보조 모델 부호 일치. "
    "H3: 게이팅 이득이 보류율과 함께 움직이며 구간이 좁을 수 있음(과잉 보류가 대가). "
    "정본 판정표(§10) 추가 행: '보류율 높음 + PR-2 유의 = 좁은 구간에서만 이득', "
    "'보류율 높음 + PR-2 비유의 = 과잉 보류가 대가만 남김'." + TAG,
)

prs.save(str(DST))
print("saved:", DST)
