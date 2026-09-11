# -*- coding: utf-8 -*-
r"""
Convert '연구계획서20260911_통합본_최종수정본.docx' into clean, beautiful GitHub Flavored Markdown.
Saves to:
  1. C:/Users/passp/OneDrive/바탕 화면/jeayong/capstone/연구계획서20260911_통합본_최종수정본.md
  2. C:/Users/passp/OneDrive/바탕 화면/jeayong/capstone/03_연구계획_및_정리노트/02_연구_및_실험계획서/2026.09.11_연구계획서_통합본_최종수정본.md
"""

import os
from pathlib import Path
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

def docx_to_markdown():
    docx_path = Path(r"C:\Users\passp\OneDrive\바탕 화면\jeayong\capstone\연구계획서20260911_통합본_최종수정본.docx")
    if not docx_path.exists():
        raise FileNotFoundError(f"File not found: {docx_path}")

    doc = docx.Document(str(docx_path))
    table_map = {t._element: t for t in doc.tables}
    
    md_lines = []
    
    for child in doc.element.body:
        if child.tag.endswith('p'):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            
            # Check font formatting of runs
            size_pt = None
            bold = False
            for r in p.runs:
                if r.font.size:
                    size_pt = r.font.size.pt
                if r.font.bold:
                    bold = True
            
            # Determine heading levels
            if size_pt and size_pt >= 15.0 and bold:
                md_lines.append(f"# {text}\n")
            elif size_pt and size_pt == 10.5 and bold and ("최종" in text or "개정" in text):
                md_lines.append(f"### {text}\n")
            elif size_pt and size_pt >= 12.0 and bold:
                md_lines.append(f"\n## {text}\n")
            elif size_pt and size_pt >= 10.5 and bold:
                md_lines.append(f"\n### {text}\n")
            elif size_pt and size_pt >= 9.5 and bold and (text.startswith("4.2.") or text.startswith("■")):
                md_lines.append(f"\n#### {text}\n")
            elif "한동대학교" in text and "Human Robotics Lab" in text:
                md_lines.append(f"{text}\n")
            else:
                # Regular paragraph
                # Check if it's bullet list or numbered list
                if text.startswith("■ ") or text.startswith("1. ") or text.startswith("2. ") or text.startswith("3. ") or text.startswith("4. ") or text.startswith("5. "):
                    md_lines.append(f"{text}\n")
                else:
                    md_lines.append(f"{text}\n")
                    
        elif child.tag.endswith('tbl'):
            if child not in table_map:
                continue
            tbl = table_map[child]
            
            # Check if 1x1 callout or code box
            if len(tbl.rows) == 1 and len(tbl.columns) == 1:
                cell = tbl.cell(0, 0)
                cell_p = cell.paragraphs
                
                # Check if code block (e.g. JSON or Consolas)
                is_code = False
                for p in cell_p:
                    for r in p.runs:
                        if r.font.name == 'Consolas':
                            is_code = True
                            break
                    if is_code:
                        break
                
                full_cell_text = "\n".join(p.text for p in cell_p if p.text.strip())
                if is_code or "{\n" in full_cell_text or "You are an expert" in full_cell_text:
                    lang = "json" if full_cell_text.strip().startswith("{") else "text"
                    md_lines.append(f"\n```{lang}\n{full_cell_text}\n```\n")
                else:
                    # Callout blockquote
                    callout_lines = []
                    for pi, p in enumerate(cell_p):
                        p_txt = p.text.strip()
                        if not p_txt:
                            continue
                        for line in p_txt.split('\n'):
                            line = line.strip()
                            if not line:
                                continue
                            if pi == 0 and line.startswith("■"):
                                callout_lines.append(f"> **{line}**")
                                callout_lines.append(">")
                            elif line.startswith("■"):
                                callout_lines.append(f"> **{line}**")
                            else:
                                callout_lines.append(f"> {line}")
                    md_lines.append("\n" + "\n".join(callout_lines) + "\n")
            
            else:
                # Standard markdown table (ensuring each row is strictly a single line)
                md_table = []
                for ri, row in enumerate(tbl.rows):
                    row_cells = []
                    for cell in row.cells:
                        # Replace newlines with <br> for strict markdown table compatibility
                        parts = []
                        for p in cell.paragraphs:
                            txt = p.text.strip()
                            if txt:
                                parts.append(txt.replace('\r\n', '<br>').replace('\n', '<br>'))
                        cell_txt = "<br>".join(parts)
                        cell_txt = cell_txt.replace("|", "&#124;")
                        row_cells.append(cell_txt)
                    md_table.append("| " + " | ".join(row_cells) + " |")
                    
                    # Add delimiter after header
                    if ri == 0:
                        delims = [":---" for _ in row_cells]
                        md_table.append("| " + " | ".join(delims) + " |")
                
                md_lines.append("\n" + "\n".join(md_table) + "\n")

    full_markdown = "\n".join(md_lines)
    
    # Clean up double blank lines
    import re
    full_markdown = re.sub(r'\n{3,}', '\n\n', full_markdown)
    
    # 1. Target in Capstone root directory
    out_file1 = Path(r"C:\Users\passp\OneDrive\바탕 화면\jeayong\capstone\연구계획서20260911_통합본_최종수정본.md")
    with open(out_file1, "w", encoding="utf-8") as f:
        f.write(full_markdown)
    print(f"Successfully generated Markdown in root: {out_file1}")

    # 2. Target in 03_연구계획_및_정리노트 directory
    out_file2 = Path(r"C:\Users\passp\OneDrive\바탕 화면\jeayong\capstone\03_연구계획_및_정리노트\02_연구_및_실험계획서\2026.09.11_연구계획서_통합본_최종수정본.md")
    out_file2.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file2, "w", encoding="utf-8") as f:
        f.write(full_markdown)
    print(f"Successfully generated Markdown in notes folder: {out_file2}")

if __name__ == '__main__':
    docx_to_markdown()
